#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
회차별 스냅샷(data/snapshots/*.json) → 시간대별 경쟁률 변화 엑셀/CSV
- 사용:
    python export_excel.py                       # 전체, 스냅샷 전부 → export/경쟁률_시간대별.xlsx + .csv
    python export_excel.py --hourly              # 시각 열을 1시간 간격으로 솎음(엑셀 가벼움)
    python export_excel.py --region 충남 대전 세종 # 지역 필터
    python export_excel.py --univ 순천향 충남대    # 대학명 부분일치 필터
    python export_excel.py --extra "D:\\기타스냅샷폴더"   # GitHub 아티팩트 등 다른 폴더의 latest.json 도 합침
- 시트:
    지원인원  : 행=대학·전형·모집단위, 열=수집시각 (지원인원)
    경쟁률    : 같은 배열, 경쟁률
    전형별    : 전형 단위 합계 시계열
    대학별    : 대학 총계 시계열
    long      : (CSV) 분석용 긴 형식 — 대학,전형,모집단위,시각,모집,지원,경쟁률
"""
import argparse
import csv
import glob
import json
import os
import re

ROOT = os.path.dirname(os.path.abspath(__file__))
SNAP = os.path.join(ROOT, "data", "snapshots")
OUT_DIR = os.path.join(ROOT, "export")


def load_snapshots(extra_dirs):
    files = sorted(glob.glob(os.path.join(SNAP, "2027_*.json")))
    for d in extra_dirs or []:
        files += sorted(glob.glob(os.path.join(d, "**", "*.json"), recursive=True))
    snaps = {}
    for f in files:
        try:
            with open(f, encoding="utf-8") as fh:
                d = json.load(fh)
        except Exception:
            continue
        if not isinstance(d, dict) or "universities" not in d:
            continue
        t = d.get("collectedAt")
        if not t:
            continue
        # 같은 시각(분 단위)이 여러 파일이면 나중 것으로
        snaps[t] = d
    return [snaps[t] for t in sorted(snaps)]


def thin_hourly(snaps):
    out, seen = [], set()
    for s in snaps:
        key = s["collectedAt"][:13]   # YYYY-MM-DDTHH
        if key in seen:
            continue
        seen.add(key)
        out.append(s)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--hourly", action="store_true")
    ap.add_argument("--region", nargs="*")
    ap.add_argument("--univ", nargs="*")
    ap.add_argument("--extra", nargs="*", help="추가 스냅샷 폴더")
    ap.add_argument("--out", default=os.path.join(OUT_DIR, "경쟁률_시간대별"))
    args = ap.parse_args()

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    unis = {u["id"]: u for u in json.load(open(os.path.join(ROOT, "universities.json"), encoding="utf-8"))}
    region_of = {u["name"]: u["region"] for u in unis.values()}
    snaps = load_snapshots(args.extra)
    if not snaps:
        raise SystemExit("스냅샷이 없습니다: " + SNAP)
    if args.hourly:
        snaps = thin_hourly(snaps)
    times = [s["collectedAt"] for s in snaps]
    print("스냅샷 %d회차 (%s ~ %s)" % (len(times), times[0], times[-1]))

    def keep(name):
        if args.region and region_of.get(name) not in args.region:
            return False
        if args.univ and not any(k in name for k in args.univ):
            return False
        return True

    # 키: (대학명, 전형, 정원구분, 모집단위) — id 는 목록 재부여로 바뀔 수 있어 이름 기준
    units, types, totals = {}, {}, {}
    quota = {}
    for ti, s in enumerate(snaps):
        for rec in s["universities"].values():
            if not rec.get("ok") or not keep(rec["name"]):
                continue
            name = rec["name"]
            t = rec.get("total") or {}
            totals.setdefault(name, {})[ti] = (t.get("quota"), t.get("app"), t.get("ratio"))
            for x in rec.get("types", []):
                k = (name, x["name"], x.get("group") or "")
                types.setdefault(k, {})[ti] = (x.get("quota"), x.get("app"), x.get("ratio"))
            for x in rec.get("units", []):
                k = (name, x["type"], x.get("group") or "", x.get("college") or "", x["unit"])
                units.setdefault(k, {})[ti] = (x.get("quota"), x.get("app"), x.get("ratio"))
                if x.get("quota") is not None:
                    quota[k] = x["quota"]
    print("모집단위 %d, 전형 %d, 대학 %d" % (len(units), len(types), len(totals)))

    os.makedirs(os.path.dirname(args.out) or ".", exist_ok=True)
    hdr_font = Font(bold=True)
    hdr_fill = PatternFill("solid", fgColor="E8EEF9")

    def write_sheet(ws, head, rows, series, idx):
        ws.append(head + times)
        for c in ws[1]:
            c.font = hdr_font
            c.fill = hdr_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for key, base in rows:
            ser = series[key]
            ws.append(list(base) + [(ser[ti][idx] if ti in ser else None) for ti in range(len(times))])
        ws.freeze_panes = ws.cell(row=2, column=len(head) + 1)
        for i in range(1, len(head) + 1):
            ws.column_dimensions[get_column_letter(i)].width = 14 if i > 1 else 18
        for i in range(len(head) + 1, len(head) + len(times) + 1):
            ws.column_dimensions[get_column_letter(i)].width = 11

    wb = Workbook()
    ws = wb.active
    ws.title = "지원인원"
    ukeys = sorted(units.keys(), key=lambda k: (region_of.get(k[0], ""), k))
    urows = [(k, (region_of.get(k[0], ""), k[0], k[1], k[2], k[3], k[4], quota.get(k))) for k in ukeys]
    uhead = ["지역", "대학", "전형", "정원", "계열", "모집단위", "모집인원"]
    write_sheet(ws, uhead, urows, units, 1)
    write_sheet(wb.create_sheet("경쟁률"), uhead, urows, units, 2)
    tkeys = sorted(types.keys(), key=lambda k: (region_of.get(k[0], ""), k))
    trows = [(k, (region_of.get(k[0], ""), k[0], k[1], k[2])) for k in tkeys]
    ws3 = wb.create_sheet("전형별")
    write_sheet(ws3, ["지역", "대학", "전형", "정원"], trows, types, 1)
    ws4 = wb.create_sheet("대학별")
    ukeys2 = sorted(totals.keys(), key=lambda n: (region_of.get(n, ""), n))
    write_sheet(ws4, ["지역", "대학"], [(n, (region_of.get(n, ""), n)) for n in ukeys2], totals, 1)
    ws5 = wb.create_sheet("안내")
    for line in [
        "2027학년도 수시모집 경쟁률 시간대별 수집 자료",
        "출처: 각 대학 입학처 경쟁률 공개 페이지(진학어플라이·유웨이어플라이 게시) 자동 수집",
        "열(시각)은 수집 시각(한국시간)이며 대학 공개 기준 시각은 이보다 5~30분 앞설 수 있음",
        "빈 칸은 해당 회차에 수집이 안 된 것(대학 미공개·수집 실패)",
        "정원외 전형 일부는 모집인원이 전형 총원으로 반복 표기됨(대학 공개 방식)",
        "제작: 충청남도교육청진로융합교육원 교육연구사 정재연",
    ]:
        ws5.append([line])
    xlsx = args.out + ".xlsx"
    wb.save(xlsx)

    csv_path = args.out + ".csv"
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["지역", "대학", "전형", "정원", "계열", "모집단위", "시각", "모집인원", "지원인원", "경쟁률"])
        for k in ukeys:
            ser = units[k]
            for ti in sorted(ser):
                q, a, r = ser[ti]
                w.writerow([region_of.get(k[0], ""), k[0], k[1], k[2], k[3], k[4], times[ti], q, a, r])
    print("저장: %s (%.1f MB), %s (%.1f MB)" % (xlsx, os.path.getsize(xlsx) / 1e6, csv_path, os.path.getsize(csv_path) / 1e6))


if __name__ == "__main__":
    main()
